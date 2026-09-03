from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .compare import filter_matrix
from .models import ComparisonResult, VersionMatrix

HEADER_FILL = PatternFill("solid", fgColor="548235")
HEADER_FONT = Font(color="FFFFFF", bold=True)
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")
RED_FONT = Font(color="C00000", bold=True)
THIN = Border(
    left=Side(style="thin", color="B8C4B0"),
    right=Side(style="thin", color="B8C4B0"),
    top=Side(style="thin", color="B8C4B0"),
    bottom=Side(style="thin", color="B8C4B0"),
)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _qty(value: float | None):
    if value is None:
        return None
    if float(value).is_integer():
        return int(value)
    return value


def _write_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def _autosize(ws: Worksheet, min_width: int = 10, max_width: int = 22) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_width
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(value) + 2))
        ws.column_dimensions[letter].width = width


def _sheet_qty_table(ws: Worksheet, result: ComparisonResult, parts) -> None:
    labels = [bom.label for bom in result.boms]
    headers = ["Codigo", "Descripcion", "Referencia", "Fabricante", *[f"qty_{label}" for label in labels]]
    _write_header_row(ws, 1, headers)
    ws.freeze_panes = "A2"
    for i, part in enumerate(parts, start=2):
        values = [
            part.key,
            part.descripcion,
            part.referencia,
            part.fabricante,
            *[_qty(part.quantities.get(label)) for label in labels],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(i, col, value)
            cell.alignment = LEFT
            cell.border = THIN
    last = max(2, len(parts) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    _autosize(ws, min_width=12, max_width=42)


def _sheet_comunes(wb: Workbook, result: ComparisonResult) -> None:
    ws = wb.active
    ws.title = "Comunes"
    _sheet_qty_table(ws, result, result.shared)


def _sheet_diferencias(wb: Workbook, result: ComparisonResult) -> None:
    ws = wb.create_sheet("Diferencias")
    _sheet_qty_table(ws, result, result.partial)


def _sheet_exclusivos(wb: Workbook, result: ComparisonResult) -> None:
    ws = wb.create_sheet("Exclusivos")
    headers = ["Codigo", "Producto", "Qty", "Descripcion", "Referencia", "Fabricante"]
    _write_header_row(ws, 1, headers)
    ws.freeze_panes = "A2"
    parts = result.exclusive
    for i, part in enumerate(parts, start=2):
        values = [
            part.key,
            part.bom,
            _qty(part.quantity),
            part.descripcion,
            part.referencia,
            part.fabricante,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(i, col, value)
            cell.alignment = LEFT
            cell.border = THIN
    last = max(2, len(parts) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    _autosize(ws, min_width=12, max_width=42)


def _sheet_matrix(wb: Workbook, matrix: VersionMatrix, title: str | None = None) -> None:
    ws = wb.create_sheet(title or "Versiones")
    n = len(matrix.versions)
    last_col = 1 + 3 * n
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    title_cell = ws.cell(1, 1, matrix.familia_corta)
    title_cell.fill = HEADER_FILL
    title_cell.font = HEADER_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_cell.border = THIN
    ws.cell(2, 1).fill = HEADER_FILL
    ws.cell(2, 1).border = THIN

    for i, version in enumerate(matrix.versions):
        start = 2 + i * 3
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + 2)
        head = ws.cell(1, start, version.label)
        head.fill = HEADER_FILL
        head.font = HEADER_FONT
        head.alignment = CENTER
        for col in range(start, start + 3):
            ws.cell(1, col).fill = HEADER_FILL
            ws.cell(1, col).border = THIN
            ws.cell(1, col).font = HEADER_FONT
        for offset, name in enumerate(["Codigo", "Layer", "Fitted"]):
            cell = ws.cell(2, start + offset, name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = THIN

    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(last_col)}2"

    for r, row in enumerate(matrix.rows, start=3):
        des = ws.cell(r, 1, row.designator)
        des.alignment = LEFT
        des.border = THIN
        for i, cell in enumerate(row.cells):
            start = 2 + i * 3
            values = [
                cell.codigo if cell.present else "",
                cell.layer if cell.present else "",
                cell.fitted if cell.present else "",
            ]
            for offset, value in enumerate(values):
                target = ws.cell(r, start + offset, value)
                target.alignment = CENTER
                target.border = THIN
                if cell.yellow:
                    target.fill = YELLOW_FILL
                if offset == 0 and cell.code_red:
                    target.font = RED_FONT

    ws.column_dimensions["A"].width = 14
    for i in range(n):
        ws.column_dimensions[get_column_letter(2 + i * 3)].width = 16
        ws.column_dimensions[get_column_letter(3 + i * 3)].width = 10
        ws.column_dimensions[get_column_letter(4 + i * 3)].width = 12
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 18


def write_excel(
    result: ComparisonResult,
    path: str | Path,
    *,
    familia: str | None = None,
    versiones: list[str] | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _sheet_comunes(wb, result)
    _sheet_diferencias(wb, result)
    _sheet_exclusivos(wb, result)

    matrices = result.version_matrices
    if familia:
        matrices = [m for m in matrices if m.familia == familia]
    written = []
    for matrix in matrices:
        filtered = filter_matrix(matrix, versiones)
        if filtered is None:
            continue
        written.append(filtered)
    if len(written) == 1:
        _sheet_matrix(wb, written[0], "Versiones")
    else:
        for matrix in written:
            _sheet_matrix(wb, matrix, matrix.familia_corta[:31])

    wb.save(path)
    return path
